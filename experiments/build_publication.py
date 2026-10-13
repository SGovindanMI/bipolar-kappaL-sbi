"""Create an integrated article and supplementary tables from verified results.

Run with the bundled document Python, after python -m src.publication_figures.
This changes presentation only; it never modifies numerical research outputs.
"""
import re,json
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from build_manuscript_docx import el,r,sub,sup,frac,equation,equations,table

ROOT=Path(__file__).resolve().parents[1];RUN=ROOT/'runs/v6_main';OUT=RUN/'publication'
LABELS={'mobility_spline':'independent-band flexible-lattice model','mobility_extra':'independent-band rising-channel model','mobility':'independent-band model','spline':'flexible-lattice model','extra':'rising-channel model','baseline':'reference model'}
CAPTIONS={
1:('figure1_network','Figure 1 | Neural posterior estimator and training histories. a, Fifty standardised transport features enter three fully connected layers, each with 256 SiLU units. Separate heads parameterise the continuous Gaussian precision and carrier-sign probability. Physical support rejection follows neural sampling; 128 accepted draws are targeted, with at least 64 required for scoring. b, Validation negative log likelihood (NLL) minus the minimum within each run, shown on a symmetric logarithmic scale. Subtracting each minimum permits visual inspection of optimisation progress; it does not make likelihood values or posterior quality comparable across models. Each curve is one training run, with 588,000 training and 12,000 validation simulations. The shared-band flexible-lattice run reached the 900-epoch cap. Source data include every plotted epoch.'),
2:('figure2_tradeoff','Figure 2 | Coverage, uncertainty width and point-estimate error on common lattice-upturn cases. a,b, Nominal 95% pointwise lattice coverage and median interval width on the 509 synthetic curves scored by all six models. c, Median absolute 400–750 K lattice-slope error on the further common subset of 87 admitted curves with positive finite endpoints. The dashed line in a marks nominal coverage. These are descriptive summaries of a fixed development benchmark; bars are not uncertainty intervals. Flexible models increase coverage and interval width without improving the median absolute slope error on the common subset.'),
3:('figure3_corpus','Figure 3 | Thermal-model dependence on the same experimental population. Points show the median per-curve difference between the slope of posterior median lattice conductivity (C) and the slope after unipolar Lorenz subtraction (B), evaluated from 400 to 750 K. All models use the same 1,528 selected curves from 649 source papers. Horizontal intervals are 95% intervals from 2,000 paper-clustered bootstrap resamples. They quantify sampling variation within this retained population and do not include uncertainty over physical-model choice. The vertical line marks zero shift.'),
4:('figure4_transfer','Figure 4 | Lattice attribution under a change in the thermal truth model. a, Independent-band rising-channel truth. b, Independent-band flexible-lattice truth. Each source contains 1,000 independently generated narrow-gap curves; labels report the number scored by each inference model. Bars show nominal 95% pointwise lattice coverage conditional on those scored cases; the dashed line marks 95%. Changing denominators preclude interpreting bar differences as paired improvements. Inference uses the same recorded seed of 7300. These simulated alternatives do not establish which mechanism is present in experimental specimens.')}

NN='''### Neural density estimation and amortisation

The network approximates the simulator posterior across many possible observations, rather than optimising each experimental curve separately. Simulator generation and network training constitute the up-front cost; the trained weights are reused for all subsequent curves. This is the meaning of amortised inference. Experimental corpus curves are not training labels. This study uses an established conditional Gaussian architecture, not a newly invented neural-network algorithm.

Inputs are the 50 features specified above. The 50-dimensional input and D continuous target coordinates are standardised using means and standard deviations computed only on the 588,000 training examples. Three fully connected hidden layers each contain 256 units with SiLU activation. Four linear heads output a D-dimensional mean, D diagonal log-precision factors, D(D−1)/2 strictly lower-triangular factors and one carrier-sign logit. D is 13, 11, 13, 14 and 16 for the independent-band, rising-channel, flexible-lattice, independent-band rising-channel and independent-band flexible-lattice variants, respectively. Continuous-head weights and biases are initialised to zero; the sign head retains its default initialisation.

The conditional density factorises between standardised continuous coordinates y and binary carrier sign s given standardised input x. With sign logit a and lower-triangular factor A, the precision is A Aᵀ, not the covariance. The diagonal entries are exp(d), with d clipped to [−3,6], and the strictly lower-triangular entries are unconstrained network outputs.

@NN1

The per-example training objective is the Gaussian negative log likelihood plus binary cross entropy; its minibatch mean is minimised. The diagonal-factor term is the log-determinant contribution. No experimental channel labels enter the loss.

@NN2

Optimisation uses AdamW with initial learning rate 0.0015, weight decay 10⁻⁵ and batch size 1,024. Gradients are clipped to norm 5. Validation is evaluated in batches of 4,096. ReduceLROnPlateau halves the learning rate after its patience of six validation epochs. A checkpoint is replaced only when validation loss improves by more than 10⁻⁴; training stops after 50 epochs without such improvement or at 900 epochs. A fixed permutation with seed 3 defines the training–validation split. Training seeds are 6101, 6111, 6121, 6131 and 6141 in the variant order above. The saved checkpoint with the best qualifying validation loss is used for inference, not necessarily the last epoch.

Network training uses single precision; the physical forward model uses double precision. Random generators are seeded, deterministic PyTorch algorithms are requested, TF32 is disabled and four CPU threads are configured. The recorded environment uses PyTorch 2.11.0 with CUDA 12.8 on an NVIDIA GeForce RTX 5060 Laptop GPU. These settings record reproducibility controls and do not imply bitwise agreement across software or hardware versions. Completed epochs and selected checkpoint epochs are reported in Supplementary Table 8. Only one training seed per variant was evaluated. Reaching the epoch cap is not a convergence certificate.

At inference, continuous proposals are sampled by a triangular solve, then transformed back to the simulator coordinates. Sign proposals are sampled independently conditional on the same input. Rejection by the prior bounds and peak-zT support produces a truncated approximate neural posterior; this is not an exact Bayesian correction of the density approximation.

@NN3

The Gaussian can encode linear covariance but cannot generally express disconnected or strongly non-Gaussian solutions. Conditional independence of sign and continuous parameters is another restriction. Validation NLL is an optimisation criterion, not proof of calibrated uncertainty, and NLL values should not be ranked across targets of different dimension. Independent known-truth coverage tests, transfer tests and selection counts therefore accompany the network. Simulation-based rank calibration, flexible-density alternatives, training-seed replication and comparison with per-curve MCMC remain outstanding; none is claimed to have been completed.
'''

def clean(t):
    t=t.replace('The released v5 checkpoint and arrays are retained as the baseline.','The corrected shared-band, monotone-lattice reference checkpoint and its inference arrays are retained as a frozen comparator; it was not retrained in this study.')
    t=t.replace('The corrected v5 release is preserved. Version 6 modules,','The corrected reference implementation (software release v5) is preserved. The extended implementation,')
    t=t.replace('Five variants are fitted: mobility, extra, spline, mobility_extra and mobility_spline.','Five variants are fitted: independent bands, a rising channel, a flexible lattice, independent bands with a rising channel, and independent bands with a flexible lattice.')
    for old in ['mobility_spline','mobility_extra','baseline']:t=re.sub(r'\b'+old+r'\b',LABELS[old],t)
    t=t.replace('reference model release','reference release')
    t=t.replace('The known-truth control from the preceding audit remains useful: 608 of 985 admitted curves showed','In the decreasing-lattice synthetic control, 608 of 985 admitted curves showed')
    t=t.replace('The contribution of this rerun','The contribution of this study')
    return t

def setup():
    d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=s.bottom_margin=Cm(2);s.left_margin=s.right_margin=Cm(2)
    for name,size in [('Normal',11),('Title',17),('Heading 1',13),('Heading 2',11.5)]:
        style=d.styles[name];style.font.name='Times New Roman';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        for node in list(style.element.xpath('.//w:pBdr')):node.getparent().remove(node)
    d.styles['Normal'].paragraph_format.line_spacing=1.12;d.styles['Normal'].paragraph_format.space_after=Pt(6)
    d.styles['Title'].font.bold=True
    for name in ['Heading 1','Heading 2']:d.styles[name].paragraph_format.keep_with_next=True
    d.settings.element.append(el('w:doNotAutoCompressPictures'))
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer._p.append(el('w:fldSimple',**{'w:instr':'PAGE'}))
    return d

def nn_equation(d,n):
    if n==1:
        equation(d,[sub('q','φ'),r('(y,s | x) = N(y; μ, '),sup([r('Λ')],[r('−1')]),r(') Bernoulli(s; sigmoid(a))')])
        equation(d,[r('Λ = A'),sup([r('A')],[r('T')]),r(',    '),sub('A','ii'),r(' = exp('),sub('d','i'),r(')')])
    elif n==2:
        equation(d,[r('ℓ = '),frac([r('1')],[r('2')]),r('‖'),sup([r('A')],[r('T')]),r('(y − μ)‖² − ∑'),sub('d','i'),r(' + '),frac([r('D')],[r('2')]),r(' ln(2π) + BCE(a,s)')])
    else:
        equation(d,[r('y = μ + '),sup([r('A')],[r('−T')]),r('ε,    ε ∼ N(0, I)')])

def render(text,path):
    doc=setup();lines=text.splitlines();i=0;eq=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line=='$$':
            equations(doc,eq);eq+=1;i+=1
            while i<len(lines) and lines[i].strip()!='$$':i+=1
        elif line.startswith('@NN'):nn_equation(doc,int(line[-1]))
        elif line.startswith('@FIG'):
            n=int(line[-1]);name,caption=CAPTIONS[n];doc.add_picture(str(OUT/'figures'/(name+'.png')),width=Cm(16.5));doc.paragraphs[-1].paragraph_format.keep_with_next=True
            p=doc.add_paragraph(caption);p.paragraph_format.line_spacing=1.05;p.paragraph_format.keep_together=True
            for run in p.runs:run.font.size=Pt(9)
        elif line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(set(x)<=set('-: ') for x in cells):rows.append(cells)
                i+=1
            table(doc,rows);continue
        elif line.startswith('# '):doc.add_paragraph(line[2:],style='Title')
        elif line.startswith('## '):doc.add_heading(line[3:],1)
        elif line.startswith('### '):doc.add_heading(line[4:],2)
        else:
            p=doc.add_paragraph(line.replace('W m^-1 K^-1','W m⁻¹ K⁻¹').replace('cm^-3','cm⁻³').replace('cm^2 V^-1 s^-1','cm² V⁻¹ s⁻¹'))
            if re.match(r'^\[\d+\]',line):
                for run in p.runs:run.font.size=Pt(9.5)
        i+=1
    # Remove all template paragraph borders, including theme-linked title decoration.
    for node in list(doc.element.xpath('.//w:pBdr')):node.getparent().remove(node)
    doc.core_properties.title=text.splitlines()[0][2:];doc.core_properties.subject='Model-sensitivity study: article-format revision, pending external validation'
    doc.save(path)

def main():
    OUT.mkdir(exist_ok=True)
    source=(RUN/'manuscript_v6_research.md').read_text(encoding='utf-8')
    intro=source.split('## 1 Introduction\n')[1].split('## 2 Methods\n')[0]
    methods=source.split('## 2 Methods\n')[1].split('## 3 Results\n')[0]
    results=source.split('## 3 Results\n')[1].split('## 4 Discussion\n')[0]
    discussion=source.split('## 4 Discussion\n')[1].split('## Data and code availability\n')[0]
    refs=source.split('## References\n')[1]
    tables=re.findall(r'(?:^\|.*\n)+',results,re.M)
    assert len(tables)==7
    for tab in tables:results=results.replace(tab,'')
    results=re.sub(r'\bTable ([1-7])\b',r'Supplementary Table \1',results)
    results=results.replace('### 3.1 Independent same-model tests','### Neural inference and independent same-model tests\n\nThe amortised estimator reuses a trained conditional density network across observations (Fig. 1). The architecture contains three hidden layers and separate continuous-density and sign heads. Full training and sampling equations are provided in Methods. The five estimators completed 178–900 epochs; optimisation histories alone do not establish uncertainty calibration.\n\n@FIG1')
    results=results.replace('### 3.4 Common-case stress comparisons','### Common-case stress comparisons')
    results=results.replace('Table 4 compares','Table 4 and Fig. 2 compare')
    anchor='Thus higher coverage does not establish generally more accurate point estimates.'
    # Figure follows the full paragraph, not a broken sentence.
    paragraphs=results.split('\n\n')
    for i,p in enumerate(paragraphs):
        if anchor in p:paragraphs[i]=p+'\n\n@FIG2'
        if p.strip().startswith('Native populations differ'):paragraphs[i]='@FIG3\n\n'+p
        if p.strip().startswith('Every estimator is also evaluated'):paragraphs[i]=p+'\n\n@FIG4'
    results='\n\n'.join(paragraphs)
    results=results.replace("The table's slope statistic", "The slope statistic in Supplementary Table 2")
    results=results.replace('Supplementary Table 3 compares','Figure 3 and Supplementary Table 3 compare')
    results=results.replace('Coverage in Supplementary Table 6','Coverage in Fig. 4 and Supplementary Table 6')
    results=results.replace('Successful training alone does not establish calibration.','The independent-band flexible-lattice model scores 5,986 of 6,000 curves and admits 5,842; lattice and bipolar coverages are 91.7% and 92.0%. Lattice coverage across the five extensions ranges from 90.9% to 94.7%, below the nominal 95% target. Successful training alone does not establish calibration.')
    results=results.replace('These tests isolate a prefactor departure','For ratios of 0.1 and 10, bipolar coverage changes from 75.5% and 74.9% in the reference model to 89.4% and 90.5% with independent bands, on 790 and 648 common scored curves, respectively. These tests isolate a prefactor departure')
    results=results.replace('The paired change subtracts the baseline shift curve by curve on this same population.','The reference median shift is −0.107 [−0.122, −0.092], compared with −0.593 [−0.702, −0.481] for independent bands with a flexible lattice and −0.227 [−0.242, −0.208] for independent bands with a rising channel. These are dimensionless slope differences, not directly measured lattice corrections. The paired change subtracts the baseline shift curve by curve on this same population.')
    results=results.replace('This transfer comparison probes attribution','When a flexible-lattice estimator is applied to rising-channel truth, the independent-band model has 74.5% lattice coverage on 999 scored curves, despite 99.3% total-conductivity coverage. In the reverse transfer, the independent-band rising-channel model scores only 268 curves and has 25.8% lattice coverage. This transfer comparison probes attribution')
    # Keep figure numbering in reading order by moving corpus subsection after common-case stress.
    start=results.index('### 3.3');end=results.index('### Common-case')
    corpus=results[start:end];results=results[:start]+results[end:]
    insert=results.index('### 3.5');results=results[:insert]+corpus+results[insert:]
    methods=methods.replace('### 2.5 Validation and estimands',NN+'\n### Validation and estimands')
    abstract='Separating lattice and bipolar heat from transport measurements depends on assumptions about carrier dynamics and thermal conduction. We compare a corrected two-band reference model with five extensions that relax shared carrier parameters, allow positive non-monotone lattice conductivity or introduce a separate rising heat channel. Each neural estimator uses 588,000 training simulations, 12,000 validation simulations and 6,000 independent test curves. On 509 commonly scored lattice-upturn cases, flexible independent-band inference increases nominal 95% lattice coverage from 41.7% to 97.2%, while widening median intervals from 0.699 to 9.983 W m⁻¹ K⁻¹. Absolute slope error does not improve on the common selected subset. Across 1,528 experimental curves from 649 papers, inferred lattice trends remain sensitive to the thermal model. Same-model undercoverage and cross-model attribution failures show that relaxed physical restrictions do not establish calibrated or unique heat-channel separation. These results identify the assumptions and external constraints needed to interpret corpus-scale neural transport inversion.'
    availability='''## Data availability

The frozen experimental corpus, sample identifiers, synthetic test data, inference arrays and figure source data are included in the accompanying research archive. Experimental curves originate from Starrydata [5]; this study does not provide a newly audited raw export. A persistent public data deposit and accession identifier have not yet been established and must be supplied before submission. Availability and licensing of the original source data must be confirmed by the authors.

## Code availability

The accompanying archive contains the reference implementation, extended physical models, neural checkpoints, training histories, seeds, environment records and analysis scripts. The reference checkpoint is the corrected software release v5; it is included as a frozen comparator, not as a newly retrained model or separate published theory. New runs are stored under runs/v6_main. Reproduction instructions and source hashes are included. A public versioned repository and persistent identifier remain to be supplied by the authors.
'''
    text='# Carrier asymmetry and thermal-model dependence in neural transport decomposition\n\n## Abstract\n\n'+abstract+'\n\n## Introduction\n'+intro+'\n## Results\n'+results+'\n## Discussion\n'+discussion+'\n## Methods\n'+methods+'\n'+availability+'\n## References\n'+refs
    text=re.sub(r'### \d+\.\d+ ', '### ',text).replace('## 5 Limitations and conclusion','### Limitations and outlook')
    text=clean(text)
    # Renumber references in first-citation order after article restructuring.
    body,reference_text=text.split('## References\n');order=[]
    for match in re.finditer(r'\[([1-7](?:,[1-7])*)\]',body):
        for x in map(int,match.group(1).split(',')):
            if x not in order:order.append(x)
    mapping={v:i+1 for i,v in enumerate(order)}
    body=re.sub(r'\[([1-7](?:,[1-7])*)\]',lambda m:'['+','.join(str(mapping[int(x)]) for x in m.group(1).split(','))+']',body)
    refdict={int(m.group(1)):m.group(2) for m in re.finditer(r'^\[(\d+)\] (.*)$',reference_text,re.M)}
    text=body+'## References\n\n'+'\n\n'.join(f'[{mapping[k]}] {refdict[k]}' for k in order)+'\n'
    (OUT/'manuscript_publication.md').write_text(text,encoding='utf-8');render(text,OUT/'manuscript_publication.docx')
    titles=['Independent same-model coverage','Native narrow-gap upturn evaluation','Experimental slope comparisons','Common-case upturn accuracy and interval width','Common-case mobility-prefactor stress','Transfer between thermal truth models','Effective-scattering ambiguity']
    supplement='# Supplementary Information\n\nCarrier asymmetry and thermal-model dependence in neural transport decomposition\n\nModel names correspond to the physical hypotheses defined in the main text. All coverage values are nominal 95% pointwise coverage conditional on scored cases. Native populations can differ; common-case comparisons use explicit intersections.\n\n'
    for i,(title,t) in enumerate(zip(titles,tables),1):supplement+=f'## Supplementary Table {i} | {title}\n\n'+t+'\n'
    supplement+='## Supplementary Table 8 | Neural training completion\n\nEpochs are one-based. The selected checkpoint is the best qualifying validation checkpoint. These are single-seed runs, not replicated convergence assessments.\n\n| Variant | Completed epochs | Selected epoch | Stop reason | Seed |\n|---|---:|---:|---|---:|\n'
    for n in ['mobility','extra','spline','mobility_extra','mobility_spline']:
        m=json.loads((RUN/n/'posterior.json').read_text());supplement+=f"| {n} | {m['epochs_completed']} | {m['best_epoch']+1} | {m['stop_reason']} | {m['seed']} |\n"
    supplement+='\nFull per-curve inference outputs, comparison masks, nested matched-draw summaries, subgroup summaries and software provenance are supplied in the repository. The prior-support adaptation and inherited development-benchmark status are described in Methods. No simulation-based calibration ranks or external validation measurements are supplied.\n'
    # Short table labels avoid needlessly tall supplementary rows.
    supplement=supplement.replace('baseline','Reference').replace('mobility_spline','Independent + flexible').replace('mobility_extra','Independent + rising')
    for old,new in [('mobility','Independent bands'),('spline','Flexible lattice'),('extra','Rising channel')]:
        supplement=re.sub(r'(?<=\| )'+old+r'(?= \|)',new,supplement)
    (OUT/'supplementary_information.md').write_text(supplement,encoding='utf-8');render(supplement,OUT/'supplementary_information.docx')
    (OUT/'article_structure.json').write_text(json.dumps({'abstract_words':len(abstract.split()),'main_figures':4,'main_tables':0,'supplementary_tables':8,'baseline_included':True,'original_release_name':'v5','neural_methods_expanded':True},indent=2))
    print(OUT)
if __name__=='__main__':main()
