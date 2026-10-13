"""Create the result-dependent Word draft with native editable OMML equations.

Run with the bundled document runtime after production/postprocessing completes.
"""
import argparse
import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def el(tag,*children,**attrs):
    x=OxmlElement(tag)
    for k,v in attrs.items():x.set(qn(k),str(v))
    for c in children:x.append(c)
    return x


def r(text):
    t=el('m:t');t.text=text
    return el('m:r',t)


def sub(base,index):return el('m:sSub',el('m:e',r(base)),el('m:sub',r(index)))
def sup(base,index):return el('m:sSup',el('m:e',*base),el('m:sup',*index))
def frac(top,bottom):return el('m:f',el('m:num',*top),el('m:den',*bottom))


def equation(doc,nodes):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(8)
    p._p.append(el('m:oMathPara',el('m:oMath',*nodes)))


def equations(doc,index):
    if index==0:
        equation(doc,[sub('μ','j'),r('(T, '),sub('η','j'),r(') = '),sub('μ','j,300'),
                      sup([r('('),frac([r('T')],[r('300 K')]),r(')')],[sub('p','j')]),
                      frac([sub('F','rⱼ+1/2'),r('('),sub('η','j'),r(')')],[sub('F','1/2'),r('('),sub('η','j'),r(')')]),
                      frac([r('Γ(3/2)')],[r('Γ('),sub('r','j'),r('+3/2)')])])
    elif index==1:
        equation(doc,[sub('κ','e'),r(' = ('),sub('L','p'),sub('σ','p'),r(' + '),sub('L','n'),sub('σ','n'),r(')T')])
        equation(doc,[sub('κ','bip'),r(' = '),frac([sub('σ','p'),sub('σ','n')],[sub('σ','p'),r(' + '),sub('σ','n')]),
                      sup([r('('),sub('S','p'),r(' − '),sub('S','n'),r(')')],[r('2')]),r('T')])
    elif index==2:
        equation(doc,[r('κ = '),sub('κ','e'),r(' + '),sub('κ','bip'),r(' + '),sub('κ','L'),r(' + '),sub('κ','extra')])
        equation(doc,[sub('κ','L'),r(' = '),sub('κ','min'),r(' + A'),sup([r('('),frac([r('300 K')],[r('T')]),r(')')],[r('ω')])])
        equation(doc,[sub('κ','extra'),r(' = B'),sup([r('('),frac([r('T')],[r('300 K')]),r(')')],[r('3')])])
    elif index==3:
        equation(doc,[sub('κ','L'),r('(T) = '),sup([r('10')],[r('g(ln T)')]),r(' W '),sup([r('m')],[r('−1')]),r(' '),sup([r('K')],[r('−1')])])
    else:raise ValueError('Unexpected equation block')


def table(doc,rows):
    if doc.paragraphs:doc.paragraphs[-1].paragraph_format.keep_with_next=True
    tab=doc.add_table(rows=1,cols=len(rows[0]));tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
    widths={5:[3.0,1.7,3.8,3.8,3.8],7:[3.1,1.65,1.85,2.2,2.2,2.0,2.1]}.get(len(rows[0]),[16/len(rows[0])]*len(rows[0]))
    # The coverage table has five columns with short numeric values.
    if len(rows[0])==5 and rows[0][1].startswith('Scored'):widths=[4.2,2.4,2.0,3.5,3.5]
    elif len(rows[0])==5 and rows[0][1]=='Truth model':widths=[3.5,3.5,2.3,3.3,3.3]
    elif len(rows[0])==5 and rows[0][1]=='Mobility ratio':widths=[4.0,2.0,2.3,3.8,3.8]
    elif len(rows[0])==5 and rows[0][1]=='Lattice coverage percent':widths=[4.0,3.0,3.0,3.0,3.0]
    props=tab._tbl.tblPr
    borders=el('w:tblBorders',*[el('w:'+side,**{'w:val':'single','w:sz':'4','w:color':'D9D9D9'}) for side in ['top','left','bottom','right','insideH','insideV']])
    props.append(borders)
    for i,values in enumerate(rows):
        cells=tab.rows[0].cells if i==0 else tab.add_row().cells
        for j,(cell,value) in enumerate(zip(cells,values)):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=value.replace('mobility_spline','Independent bands + spline').replace('mobility_extra','Independent bands + extra').replace('mobility','Independent bands')
            pr=cell._tc.get_or_add_tcPr()
            pr.append(el('w:tcMar',*[el('w:'+side,**{'w:w':'90','w:type':'dxa'}) for side in ['top','left','bottom','right']]))
            if i==0:pr.append(el('w:shd',**{'w:fill':'24475B'}))
            elif i%2==0:pr.append(el('w:shd',**{'w:fill':'F3F6F8'}))
            for p in cell.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
                p.paragraph_format.keep_with_next=i==0
                for run in p.runs:
                    run.font.name='Times New Roman';run.font.size=Pt(9)
                    if i==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
        if i==0:tab.rows[0]._tr.get_or_add_trPr().append(el('w:tblHeader'))
        tab.rows[i]._tr.get_or_add_trPr().append(el('w:cantSplit'))
    for column,width in zip(tab.columns,widths):column.width=Cm(width)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)


def main():
    ap=argparse.ArgumentParser();ap.add_argument('run',type=Path);a=ap.parse_args()
    state=json.loads((a.run/'status.json').read_text())
    if state['status']!='complete':raise RuntimeError('Production is incomplete')
    source=a.run/'manuscript_v6_research.md'
    lines=source.read_text(encoding='utf-8').splitlines();doc=Document();sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        style=doc.styles[name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
    normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(6)
    doc.styles['Title'].font.size=Pt(17);doc.styles['Title'].font.bold=True
    for name,size in [('Heading 1',13),('Heading 2',11.5),('Heading 3',11)]:
        s=doc.styles[name];s.font.size=Pt(size);s.font.bold=True;s.paragraph_format.keep_with_next=True
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    foot._p.append(el('w:fldSimple',**{'w:instr':'PAGE'}))
    i=0;eq=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line=='$$':
            equations(doc,eq);eq+=1;i+=1
            while i<len(lines) and lines[i].strip()!='$$':i+=1
        elif line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(set(c)<=set('-: ') for c in cells):rows.append(cells)
                i+=1
            table(doc,rows);continue
        elif line.startswith('# '):doc.add_paragraph(line[2:],style='Title')
        elif line.startswith('## '):doc.add_heading(line[3:],level=1)
        elif line.startswith('### '):doc.add_heading(line[4:],level=2)
        elif line.startswith('[') and '] ' in line[:6]:
            p=doc.add_paragraph(line);p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(3)
            for run in p.runs:run.font.size=Pt(9.5)
        else:doc.add_paragraph(line.replace('W m^-1 K^-1','W m⁻¹ K⁻¹').replace('cm^-3','cm⁻³').replace('cm^2 V^-1 s^-1','cm² V⁻¹ s⁻¹').replace('log10','log₁₀'))
        i+=1
    if eq!=4:raise ValueError('Unexpected equation count')
    figdir=a.run/'figures'
    if (figdir/'coverage.png').exists():
        doc.add_heading('Figures',level=1)
        doc.add_picture(str(figdir/'coverage.png'),width=Cm(16.5))
        doc.paragraphs[-1].paragraph_format.keep_with_next=True
        doc.add_paragraph('Figure 1. Nominal 95% pointwise channel coverage on scored curves. Same-model tests and the shared narrow-gap lattice-upturn stress test use different truth distributions.')
        doc.add_page_break()
        doc.add_picture(str(figdir/'common_population.png'),width=Cm(16.5))
        doc.paragraphs[-1].paragraph_format.keep_with_next=True
        doc.add_paragraph('Figure 2. Median per-curve C-minus-B slope shift on the same selected experimental curves across models. Intervals resample source papers and do not include model-choice uncertainty.')
    doc.core_properties.title='Carrier asymmetry and lattice flexibility in thermal transport decomposition'
    doc.core_properties.subject='Research draft from the completed version 6 rerun'
    doc.save(a.run/'manuscript_v6_research.docx')
    print(a.run/'manuscript_v6_research.docx')


if __name__=='__main__':main()
