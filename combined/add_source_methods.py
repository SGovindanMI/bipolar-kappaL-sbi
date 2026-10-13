"""Add source-verification and ancillary-data methods to the finalised documents.

Run after build_manuscript.js and finalize_documents.py. All empirical analyses
already use the corrected corpus; no old-versus-new comparison is added here.
"""
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'combined/manuscript'
A=json.loads((ROOT/'data/ancillary_summary.json').read_text(encoding='utf-8'))
N=json.loads((ROOT/'data/manuscript_numbers.json').read_text(encoding='utf-8'))

def table(d,number,caption,headers,rows,widths):
    p=d.add_paragraph(f'Supplementary Table {number}. '+caption);p.paragraph_format.keep_with_next=True
    t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
    for j,w in enumerate(widths):t.columns[j].width=Inches(w)
    for c,v in zip(t.rows[0].cells,headers):c.text=v
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for i,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j])
            tcpr=c._tc.get_or_add_tcPr();va=OxmlElement('w:vAlign');va.set(qn('w:val'),'center');tcpr.append(va)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');tcpr.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1
                if i==0:p.paragraph_format.keep_with_next=True
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(9);r.bold=i==0
    return t

def equation(d,text):
    p=d.add_paragraph();p.paragraph_format.space_after=Pt(6)
    m=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);m.append(r);p._p.append(m)

def main():
    d=Document(OUT/'combined_manuscript.docx');s=Document(OUT/'combined_supplementary_information.docx')
    assert len(s.tables)==30,'Rebuild and finalise both documents before this step.'
    anchor=next(p for p in d.paragraphs if p.text=='3. Results')
    p=anchor.insert_paragraph_before('2.9 Source consistency and ancillary transport data',style=next(p.style for p in d.paragraphs if p.text.startswith('2.8 ')))
    anchor.insert_paragraph_before('Source verification is applied before model-specific screening. The final observation matrix and each model’s retained posterior arrays share the same row order, with immutable file hashes and source-paper/sample identities. Accepted records reproduce all three measured channels within the tolerances in Section 2.1; the quality criteria and reproduction conventions are summarised in Supplementary Table 31. The cohort is not a random sample of materials, and verification of digitised values does not establish experimental accuracy or validate inherited family labels.')
    anchor.insert_paragraph_before('We also join auxiliary curves by source paper and sample. Supplementary Table 32 reports their availability in the analysed cohort. Power factor and zT provide internal consistency checks against S²σ and S²σT/κ, respectively. Comparisons require at least five overlapping grid points and canonical SI units. For multiple auxiliary curves, all candidates are retained and the summary uses the smallest per-curve median symmetric relative difference. A difference above 20% is a descriptive review flag, not an automatic exclusion. These derived quantities reuse the fitted transport measurements and do not provide independent validation of the thermal decomposition (Supplementary Table 33).')
    anchor=next(p for p in d.paragraphs if p.text=='4. Discussion')
    anchor.insert_paragraph_before('3.12 Ancillary transport consistency',style=next(p.style for p in d.paragraphs if p.text.startswith('2.8 ')))
    pf,zt=A['Power factor'],A['ZT']
    anchor.insert_paragraph_before(f'Power-factor consistency can be evaluated for {pf["n"]:,} records and zT consistency for {zt["n"]:,}. Median per-record symmetric discrepancies are {100*pf["median_difference"]:.2f}% and {100*zt["median_difference"]:.2f}%, respectively. Even the best matching candidate exceeds the 20% review threshold for {pf["flag_n"]:,} power-factor records and {zt["flag_n"]:,} zT records; these groups may overlap. Digitisation, differing experimental conditions, source-curve selection and internally inconsistent measurements can all contribute. These flags identify records for review and are not estimates of the rate of erroneous experiments (Supplementary Table 33).')
    anchor.insert_paragraph_before('The analysed cohort also contains source-matched carrier-concentration, carrier-mobility and Hall-coefficient curves for 908, 714 and 101 specimens, respectively (Supplementary Table 32). Availability does not guarantee usable temperature overlap or comparable measurement conditions. Hall-derived concentration is not automatically the model’s net dopant density, and a two-carrier Hall comparison requires additional assumptions about Hall factors and scattering. These measurements are catalogued for subsequent validation and do not enter the present inference. Source-reported lattice or electronic thermal conductivity is likewise not treated as independent ground truth, because it may already depend on a subtraction model.')
    # Preserve the scientific need for source-paper interpretation without implying no raw export audit.
    for p in d.paragraphs:
        if p.text.startswith('A source-paper audit must record'):
            p.text='Before claiming errors in published lattice values, each source paper must still be checked for its reported quantity and thermal correction. The source-value and identity checks in Section 2.1 do not establish those methodological choices. These decomposition choices do not change measured zT; they change its interpretation.'
    heading=s.add_paragraph('Source consistency and ancillary measurements');heading.runs[0].bold=True;heading.paragraph_format.keep_with_next=True
    s.add_paragraph(f'The empirical analyses in this manuscript use {N["n"]:,} records from {N["papers"]:,} source papers. Source checks are evaluated against the ThermoelectricMaterials export in dated release data-20261008, with database snapshot 9 October 2026. The archived selection list defines the cohort. File checksums, composite identities and reproducible analysis commands are supplied with the repository; the complete historical selection rule is not recoverable. All main-text and supplementary empirical results use this same cohort before their stated model-specific masks.')
    table(s,31,'Source-verification criteria applied to the analysed cohort. Agreement with an export is not independent validation of the source experiment.',
          ['Check','Criterion'],[
              ['Identity','Source DOI plus sample ID maps to one paper SID; cross-paper integer-ID joins are disallowed.'],
              ['Properties and units','Canonical Seebeck, conductivity/resistivity and total thermal-conductivity labels with SI units.'],
              ['Temperatures','Finite usable curves; identical repeated points may collapse; conflicting repeated temperatures require review.'],
              ['Window','At least 200 K; lower endpoint 300–500 K; upper endpoint ≤1250 K; source coverage with ≤0.0001 K endpoint-rounding allowance.'],
              ['Numerical agreement','Maximum discrepancy ≤1% for σ and κ; absolute floor 0.001 S/m for σ; maximum absolute Seebeck discrepancy ≤1 µV/K.'],
              ['Inference alignment','Each model uses the same ordered observations. Retained posterior draws correspond to those observations; no auxiliary channel labels enter neural training.'],
              ['Limits','A list-defined cohort with incomplete historical selection provenance; source consistency does not establish physical correctness or unbiased material representation.']],[1.2,5.0])
    names=['Power factor','ZT','Carrier concentration','Carrier mobility','Hall coefficient','Hall mobility','Lattice thermal conductivity','Electronic thermal conductivity','Thermal diffusivity','Lorenz number','Density','Specific heat capacity (J/gK)']
    rows=[]
    for prop in names:
        records=[v for v in A['inventory'] if v['property']==prop]
        rows.append([prop,sum(v['samples'] for v in records),sum(v['temperature_K_samples'] for v in records)])
    table(s,32,'Auxiliary-property availability among the analysed records. Counts are matched specimens, not usable validation cases. Canonical Hall-mobility unit groups have disjoint specimens in this export.',
          ['Property','Matched specimens','Temperature in K'],rows,[3.4,1.4,1.4])
    s.add_paragraph('Available descriptors include material form, fabrication process, measurement method, relative density, grain size and inherited family tags; coverage is incomplete. Source-reported heat components may be model-derived. Ambiguous density and heat-capacity units preclude automatic thermal reconstruction. No auxiliary measurement enters the likelihood or establishes phase purity or independent replication. Two sample IDs have identical feature vectors and share one paper-bootstrap cluster.')
    s.add_paragraph('For auxiliary value y and transport reconstruction ŷ, take the median symmetric discrepancy over at least five overlapping points, then the minimum across candidate curves for each specimen. This optimistic choice is an internal consistency diagnostic, not independent validation.')
    equation(s,'PF = S²σ     ;     zT = S²σT/κ')
    equation(s,'d = 2|y − ŷ| / (|y| + |ŷ|)')
    table(s,33,'Internal consistency with auxiliary power-factor and zT curves. Percentages are symmetric discrepancies, not posterior coverage or experimental error rates. A specimen may be flagged for both properties.',
          ['Property','Specimens','Median discrepancy','Above 20%'],[[name,f'{v["n"]:,}',f'{100*v["median_difference"]:.2f}%',f'{v["flag_n"]:,}'] for name,v in [('Power factor',pf),('zT',zt)]],[2.0,1.3,1.6,1.3])
    for p in d.paragraphs:
        if p.text.startswith("Source measurements are from Starrydata"):p.alignment=0
        if p.text.startswith("[") and len(p.text)>30:p.paragraph_format.keep_together=True
    for doc in [d,s]:
        doc.core_properties.subject='Source-consistent thermoelectric corpus with reference and extended model analyses'
        doc.settings.element.append(OxmlElement('w:doNotAutoCompressPictures'))
    d.save(OUT/'combined_manuscript.docx');s.save(OUT/'combined_supplementary_information.docx')
    structure_path=OUT/'structure.json'
    structure=json.loads(structure_path.read_text(encoding='utf-8'))
    structure['supplementary_tables']=len(s.tables)
    structure['paragraph_words_excluding_tables']=sum(len(p.text.split()) for p in d.paragraphs)
    structure_path.write_text(json.dumps(structure,indent=2),encoding='utf-8')
    print('Added source and ancillary methods; 33 supplementary tables; no historical dataset comparison in paper.')
if __name__=='__main__':main()
