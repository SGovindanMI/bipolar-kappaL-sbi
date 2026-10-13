"""Add claim-linked citations and archived computing details after add_source_methods.

This final authoring step changes no scientific arrays. Rebuild the base documents
before repeating it. Citation numbers follow first appearance in the main text.
"""
import json,re
from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.shared import Pt

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'combined/manuscript'
REG=ROOT/'combined/reference_registry.json'

def replace(p,old,new):
    """Replace text across runs while preserving all unrelated equation and style nodes."""
    nodes=p._p.xpath('.//w:t');text=''.join(n.text or '' for n in nodes)
    assert old in text,old
    start=text.index(old);end=start+len(old);pos=0;done=False
    for n in nodes:
        value=n.text or '';a,b=pos,pos+len(value);pos=b
        if b<=start or a>=end:continue
        left=value[:max(0,start-a)];right=value[max(0,end-a):]
        n.text=left+(new if not done else '')+right;done=True

def main():
    d=Document(OUT/'combined_manuscript.docx');s=Document(OUT/'combined_supplementary_information.docx')
    assert 'Computational environment and reproducibility' not in '\n'.join(p.text for p in d.paragraphs),'Rebuild documents first.'
    refs=[p for p in d.paragraphs if re.match(r'^\[\d+\]\s',p.text)]
    assert len(refs)==12
    registry={str(i):{'text':re.sub(r'^\[\d+\]\s*','',p.text)} for i,p in enumerate(refs,1)}
    old_sources={
      '1':('https://thesis.caltech.edu/9780/1/Kim_Hyun-Sik_2016_Thesis_Final.pdf','Author thesis, chapter 8 reproduces the article; publisher endpoint inaccessible.'),
      '2':('https://pmc.ncbi.nlm.nih.gov/articles/PMC5551806/','Full text verified from user-supplied PDF (9 pages). Pages 3-5, Eq. (5), Table 1 and Fig. 4 support model-assisted bipolar decomposition; page 8 records Hall measurements. Not independent component ground truth.'),
      '3':('https://www.nature.com/articles/srep10136','Publisher full text; bipolar expression and minority-conductivity dependence.'),
      '4':('https://pubs.rsc.org/en/content/getauthorversionpdf/c9ee03921j','Author manuscript accessible; grain-boundary subtraction bias.'),
      '5':('https://www.nature.com/articles/s41467-025-66578-y','Publisher full text; low-temperature high-field separation and inelastic-scattering interpretation.'),
      '6':('https://proceedings.neurips.cc/paper/2016/hash/6aca97005c68f1206823815f66102863-Abstract.html','Proceedings abstract and linked full paper; methodological precedent, not the exact implemented architecture.'),
      '7':('https://arxiv.org/html/1911.01429v3','Author full text and journal DOI verified; methodological overview.'),
      '8':('https://arxiv.org/abs/1804.06788','Author preprint; cited as SBC methodology, not as a claim that full SBC was performed.'),
      '9':('https://www.nature.com/articles/s41467-024-51377-8','Publisher full text; calculated lattice upturn and experimental comparison in cubic GeTe.'),
      '10':('https://www.tandfonline.com/doi/full/10.1080/27660400.2025.2506976','Publisher full text; database design and digitisation. Current cohort counts are this project’s measurements.'),
      '11':('https://anthonylloyd.github.io/public/interp/pchip/A_method_for_constructing_local_monotone_picewise_cubic_interpolants_fritsch1984.pdf','Publisher metadata verified; complete published five-page article accessed through a PDF mirror. Supports local shape-preserving interpolation, including piecewise-monotone data.'),
      '12':('https://link.springer.com/article/10.1007/s11664-999-0211-y','Full text verified from user-supplied PDF (4 pages; journal pages 869-872). Approximate Seebeck-maximum gap relation; pages 870-871 examine mobility, effective-mass and scattering dependence and narrow-gap statistics. Not an independent validation using the same Seebeck data.')}
    for k,(url,access) in old_sources.items():registry[k].update(url=url,access=access)
    registry['6']['text']='Papamakarios, G. & Murray, I. Fast ε-free inference of simulation models with Bayesian conditional density estimation. Advances in Neural Information Processing Systems 29, 1028–1036 (2016). https://proceedings.neurips.cc/paper/2016/hash/6aca97005c68f1206823815f66102863-Abstract.html'
    new=[
      ('lorenz','Putatunda, A. & Singh, D. J. Lorenz number in relation to estimates based on the Seebeck coefficient. Materials Today Physics 8, 49–55 (2019). doi:10.1016/j.mtphys.2019.01.001','https://www.sciencedirect.com/science/article/abs/pii/S2542529318301561','Full text verified from user-supplied PDF (7 pages; journal pages 49-55). Equations (4) and (7) distinguish open-circuit conductivity; pages 51-54 and Table 2 support scattering- and band-structure-dependent departures from the Kim estimate. No universal error bound or corpus-specific correction follows.'),
      ('multiband','Thesberg, M., Kosina, H. & Neophytou, N. On the Lorenz number of multiband materials. Physical Review B 95, 125206 (2017). doi:10.1103/PhysRevB.95.125206','https://www.iue.tuwien.ac.at/pdf/ib_2017/JB2017_Thesberg_1.pdf','Author-hosted full paper; multiband Lorenz-number variation.'),
      ('trust','Hermans, J., Delaunoy, A., Rozet, F., Wehenkel, A., Begy, V. & Louppe, G. A trust crisis in simulation-based inference? Your posterior approximations can be unfaithful. Transactions on Machine Learning Research (2022). https://openreview.net/forum?id=LHAbHkt6Aq','https://arxiv.org/html/2110.06581v3','Author full text and published TMLR PDF; posterior overconfidence and coverage diagnostics.'),
      ('benchmark','Lueckmann, J.-M., Boelts, J., Greenberg, D., Gonçalves, P. & Macke, J. Benchmarking simulation-based inference. Proceedings of Machine Learning Research 130, 343–351 (2021). https://proceedings.mlr.press/v130/lueckmann21a.html','https://proceedings.mlr.press/v130/lueckmann21a/lueckmann21a.pdf','Proceedings full paper; complementary diagnostics and no uniformly best algorithm.'),
      ('silu','Elfwing, S., Uchibe, E. & Doya, K. Sigmoid-weighted linear units for neural network function approximation in reinforcement learning. Neural Networks 107, 3–11 (2018). doi:10.1016/j.neunet.2017.12.012','https://arxiv.org/pdf/1702.03118','Author full paper; journal metadata verified on publisher page. Citation credits the activation only.'),
      ('adamw','Loshchilov, I. & Hutter, F. Decoupled weight decay regularization. International Conference on Learning Representations (2019). https://openreview.net/forum?id=Bkg6RiCqY7','https://arxiv.org/pdf/1711.05101','Author full paper; AdamW method. Hyperparameters are this project’s recorded choices.'),
      ('torch','Paszke, A. et al. PyTorch: an imperative style, high-performance deep learning library. Advances in Neural Information Processing Systems 32 (2019). https://arxiv.org/abs/1912.01703','https://arxiv.org/abs/1912.01703','Author paper record; framework attribution. Run versions come from local manifests, not this paper.'),
      ('numpy','Harris, C. R. et al. Array programming with NumPy. Nature 585, 357–362 (2020). doi:10.1038/s41586-020-2649-2','https://www.nature.com/articles/s41586-020-2649-2','Publisher full text; software attribution.'),
      ('scipy','Virtanen, P. et al. SciPy 1.0: fundamental algorithms for scientific computing in Python. Nature Methods 17, 261–272 (2020). doi:10.1038/s41592-019-0686-2','https://www.nature.com/articles/s41592-019-0686-2','Publisher full text; software attribution.'),
      ('gap','Gibbs, Z. M., Kim, H.-S., Wang, H. & Snyder, G. J. Band gap estimation from temperature dependent Seebeck measurement—Deviations from the 2e|S|maxTmax relation. Applied Physics Letters 106, 022112 (2015). doi:10.1063/1.4905922','https://thesis.caltech.edu/9780/1/Kim_Hyun-Sik_2016_Thesis_Final.pdf','Author thesis chapter 7 reproduces the article; author institutional record verifies DOI. Supports mobility-ratio and narrow-gap limitations.')]
    for key,text,url,access in new:registry[key]={'text':text,'url':url,'access':access}
    # Mark existing citations with keys before inserting any new citations.
    refstart=next(i for i,p in enumerate(d.paragraphs) if p.text=='References')
    for p in d.paragraphs[:refstart]:
        for node in p._p.xpath('.//w:t'):
            if node.text:node.text=re.sub(r'\[(\d+(?:,\d+)*)\]',lambda m:'[@'+m.group(1)+']',node.text)
    def find(prefix):return next(p for p in d.paragraphs if p.text.startswith(prefix))
    p=find('The figure of merit');replace(p,'Its applicability depends on the assumed carrier and scattering regime.','Its applicability depends on the assumed carrier and scattering regime; realistic non-parabolic band structures can produce departures from this approximation [@lorenz].')
    p=find('With two carrier types');replace(p,'using signed band Seebeck coefficients.','using signed band Seebeck coefficients [@3].');replace(p,'Specifying the Lorenz convention is therefore essential.','Specifying the Lorenz convention is therefore essential, especially when comparing multiband Lorenz numbers [@multiband].')
    p=find('Amortised simulation-based');replace(p,'Its uncertainty must be checked against held-out simulations [@8].','Its uncertainty must be checked against held-out simulations [@8,trust].')
    p=find('The simulator is a two-band');replace(p,'The simulator is a two-band single-parabolic-band model.','The simulator is a two-band single-parabolic-band model, following the transport framework used for bipolar conduction [@3].')
    p=find('Ten continuous parameters');replace(p,'Bounds are inherited from the supplied study, rather than independently selected on a new real-data corpus.','Bounds were fixed during this project’s model development, rather than independently selected on a new real-data corpus.')
    p=find('The network has');replace(p,'256 SiLU units.','256 SiLU units [@silu].')
    p=find('Training uses AdamW');replace(p,'Training uses AdamW,','Training uses AdamW [@adamw],')
    p=find('We also assess empirical coverage');replace(p,'These finite-sample checks are not full joint simulation-based rank calibration.','These finite-sample checks are not full joint simulation-based rank calibration [@8]; coverage and other posterior-quality metrics assess different properties [@trust,benchmark].')
    p=find('Five extended models');replace(p,'Five extended models relax the reference restrictions.','Five extended models relax the reference restrictions. The carrier extensions are motivated by the sensitivity of bipolar transport to the separate carrier conductivities and mobilities [@3].')
    p=find('The inferred gap is not');replace(p,'reuse Seebeck data already seen by the network.','reuse Seebeck data already seen by the network and can deviate with unequal weighted mobilities or narrow gaps [@gap].')
    p=find('Further checks should examine');replace(p,'larger draw counts.','larger draw counts [@8,trust,benchmark].')
    env=json.loads((ROOT/'runs/v6_main/manifest.json').read_text())['environment']
    for variant in ['mobility','extra','spline','mobility_extra','mobility_spline']:
        assert json.loads((ROOT/f'runs/v6_main/{variant}/manifest.json').read_text())['environment']==env
    refenv=json.loads((ROOT/'data/posterior_v5.json').read_text())
    assert refenv['torch']==env['torch'] and refenv['device']=='cuda'
    anchor=find('2.5 Calibration')
    p=anchor.insert_paragraph_before('Computational environment and reproducibility');p.runs[0].bold=True;p.paragraph_format.keep_with_next=True
    anchor.insert_paragraph_before(f"The archived extended-model runs used Python {env['python']}, NumPy {env['numpy']} [@numpy], SciPy {env['scipy']} [@scipy] and PyTorch {env['torch']} [@torch], with CUDA {env['cuda']} on an {env['gpu']}. Reference-training metadata independently record CUDA execution under the same PyTorch build. Training uses single precision, while the physical forward calculations use double precision. The training code seeds its random generators, requests deterministic PyTorch algorithms, disables TF32 and configures four CPU threads. These controls do not guarantee bitwise equality across environments. Supplementary computational details distinguish recorded specifications from unrecorded fields.")
    # Separate supplementary narrative: software/environment, no additional numerical result table.
    p=s.add_paragraph('Supplementary computational details');p.runs[0].bold=True;p.paragraph_format.page_break_before=True;p.paragraph_format.keep_with_next=True
    s.add_paragraph(f"All five extended-run manifests record Python {env['python']}, NumPy {env['numpy']} [@numpy], SciPy {env['scipy']} [@scipy], PyTorch {env['torch']} [@torch], CUDA {env['cuda']} and an {env['gpu']}. Reference-training metadata record the same PyTorch build and CUDA execution, but do not independently record the GPU model. CPU model, installed RAM, operating-system version, GPU memory and driver version were not captured in the archived run manifests. They are not inferred from the present computer.")
    s.add_paragraph('The training implementations configure four CPU threads, request deterministic algorithms, disable TF32 for matrix multiplication and cuDNN, and set a default cuBLAS workspace configuration of :4096:8. Network inputs and parameters use single precision; transport evaluation promotes parameters to double precision. The archived checkpoint and inference arrays define the reported results. Software or device changes can still change numerical trajectories. The corrected empirical cohort retains unchanged inputs and matched draws, so the source correction did not require neural retraining.')
    s.add_paragraph('The SiLU activation [@silu] and AdamW optimiser [@adamw] are established components; network width, depth, precision-factor output and optimisation settings are choices of this study. Supplementary Figure 1 and Tables 1 and 17 describe training. Marginal and pointwise coverage are distinguished from full rank-based simulation calibration [@8], and no citation is used to imply that an unperformed calibration or architecture comparison has been completed [@trust,benchmark].')
    # Renumber by first occurrence. Do not match intervals or mathematical notation.
    order=[]
    for p in d.paragraphs[:next(i for i,p in enumerate(d.paragraphs) if p.text=='References')]:
        for m in re.finditer(r'\[@([^]]+)\]',p.text):
            for key in m.group(1).split(','):
                assert key in registry,key
                if key not in order:order.append(key)
    assert set(order)==set(registry),(set(registry)-set(order))
    numbering={key:i+1 for i,key in enumerate(order)}
    locations=[]
    for name,doc in [('main',d),('supplement',s)]:
        section='Opening'
        for p in doc.paragraphs:
            if re.match(r'^\d+\.',p.text) or p.text.startswith('Supplementary computational'):section=p.text
            keys=[key for m in re.finditer(r'\[@([^]]+)\]',p.text) for key in m.group(1).split(',')]
            if keys:locations.append({'document':name,'section':section,'keys':keys,'paragraph':p.text})
            for node in p._p.xpath('.//w:t'):
                if node.text:node.text=re.sub(r'\[@([^]]+)\]',lambda m:'['+','.join(str(n) for n in sorted({numbering[k] for k in m.group(1).split(',')}))+']',node.text)
    template=deepcopy(refs[0]._p)
    for p in refs:d.element.body.remove(p._p)
    def bibliography(doc,keys):
        for key in keys:
            p=doc.add_paragraph(style=refs[0].style)
            if template.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr') is not None:
                old=p._p.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr')
                if old is not None:p._p.remove(old)
                p._p.insert(0,deepcopy(template.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr')))
            r=p.add_run(f'[{numbering[key]}] '+registry[key]['text']);r.font.name='Times New Roman';r.font.size=Pt(9)
            p.paragraph_format.keep_together=True
    bibliography(d,order)
    used={k for l in locations if l['document']=='supplement' for k in l['keys']}
    p=s.add_paragraph('References cited in Supplementary Information');p.runs[0].bold=True;p.paragraph_format.keep_with_next=True
    s.add_paragraph('Reference numbers follow the main manuscript.')
    bibliography(s,[k for k in order if k in used])
    # Match the main-text paragraph spacing and type size for the added methods.
    source=find('The full precision factor')
    for p in d.paragraphs:
        if p.text.startswith('The archived extended-model runs used'):
            old=p._p.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr')
            if old is not None:p._p.remove(old)
            p._p.insert(0,deepcopy(source._p.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr')))
            for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(11)
        if p.text=='Computational environment and reproducibility':
            p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(4)
            for run in p.runs:run.font.size=Pt(11)
    for p in s.paragraphs:
        if p.text.startswith(('All five extended-run','The training implementations','The SiLU activation')):
            p.paragraph_format.space_after=Pt(6)
    for key in registry:registry[key]['number']=numbering[key]
    REG.write_text(json.dumps({'checked_on':'2026-10-10','references':registry,'placements':locations},ensure_ascii=False,indent=2),encoding='utf-8')
    (ROOT/'data/computational_environment.json').write_text(json.dumps({'extended':env,'reference_recorded':{'device':refenv['device'],'torch':refenv['torch']},'unrecorded':['CPU model','RAM','OS version','GPU memory','driver version'],'sources':['runs/v6_main/manifest.json','runs/v6_main/*/manifest.json','data/posterior_v5.json','src/posterior.py','src/relaxed_posterior.py']},indent=2))
    d.save(OUT/'combined_manuscript.docx');s.save(OUT/'combined_supplementary_information.docx')
    structure=json.loads((OUT/'structure.json').read_text());structure.update(reference_count=len(order),reference_order=order)
    (OUT/'structure.json').write_text(json.dumps(structure,indent=2))
    print(f'Added {len(order)} numbered references with paragraph-level placements and verified run details.')

if __name__=='__main__':main()
