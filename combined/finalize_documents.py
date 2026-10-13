"""Keep all analyses while moving detailed tables to SI and checking labels.

Run immediately after build_manuscript.js. It is not an in-place incremental editor.
"""
import json,re
from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'combined/manuscript'
N=json.loads((ROOT/'combined/combined_numbers.json').read_text(encoding='utf-8'))

def main():
    mainpath=OUT/'combined_manuscript.docx';suppath=OUT/'combined_supplementary_information.docx'
    doc=Document(mainpath);si=Document(suppath)
    si.styles.element.append(deepcopy(Document().styles['Table Grid'].element))
    assert len(doc.tables)==12 and len(si.tables)==15,'Rebuild with build_manuscript.js before finalising'
    moved={1:16,2:17,3:18,4:19,6:20,7:21,8:22,9:23,10:24,12:25};kept={5:1,11:2}
    # Save table/caption pairs before changing references.
    for old,tab in enumerate(list(doc.tables),1):
        if old not in moved:continue
        caption=tab._tbl.getprevious()
        assert ''.join(caption.xpath('.//w:t/text()')).startswith(f'Table {old}.')
        doc.element.body.remove(caption);doc.element.body.remove(tab._tbl)
        si.element.body.insert(len(si.element.body)-1,caption);si.element.body.insert(len(si.element.body)-1,tab._tbl)
    def replace_ref(m):
        n=int(m.group(1))
        return f'Supplementary Table {moved[n]}' if n in moved else f'Table {kept[n]}'
    for d in [doc,si]:
        for node in d.element.xpath('//w:t'):
            if node.text:node.text=re.sub(r'(?<!Supplementary )\bTable (1[0-2]|[1-9])\b',replace_ref,node.text)
    # Added analysis uses inherited family labels and existing 80-curve rule.
    p=si.add_paragraph('Supplementary Table 26. Family representation in the common six-model population. Counts include every inherited label. Named families and the heterogeneous residual other group are distinguished; effects are reported only for groups with at least 80 curves.')
    p.paragraph_format.keep_with_next=True
    t=si.add_table(rows=1,cols=4);t.style='Table Grid'
    for c,v in zip(t.rows[0].cells,['Inherited family label','Common curves','Source papers','Effect reported']):c.text=v
    for f,g in N['family_common']['families'].items():
        for c,v in zip(t.add_row().cells,[f,str(g['n']),str(g['papers']),'Yes' if g['reported'] else 'No: below 80']):c.text=v
    p=si.add_paragraph('Supplementary Table 27. Family-conditioned C − B shifts and paired changes from the reference model, on identical curves across all models. Counts are in Supplementary Table 26. Values are medians [95% paper-clustered intervals], using 2,000 bootstrap resamples with seed 2024. A paired median difference need not equal the difference of two medians. Intervals are descriptive and unadjusted for multiple comparisons.')
    p.paragraph_format.keep_with_next=True
    t=si.add_table(rows=1,cols=4);t.style='Table Grid'
    for c,v in zip(t.rows[0].cells,['Family','Model','C − B','Paired change']):c.text=v
    def fmt(v):return f'{v[0]:.3f} [{v[1]:.3f}, {v[2]:.3f}]'.replace('-','−')
    for f,g in N['family_common']['families'].items():
        if not g['reported']:continue
        for model in N['models']:
            v=g['models'][model]
            for c,text in zip(t.add_row().cells,[f,N['short'][model],fmt(v['C_B']),fmt(v['paired_change_from_reference'])]):c.text=text
    component=json.loads((ROOT/'combined/component_comparison.json').read_text(encoding='utf-8'))
    p=si.add_paragraph('Model-specific reconstruction of the thermal estimates');p.runs[0].bold=True;p.paragraph_format.keep_with_next=True
    si.add_paragraph(f"All following comparisons use the identical {component['n']:,} curves from {component['papers']:,} papers. A uses measured inputs only. B uses each model\'s own posterior median unipolar Lorenz number; C uses its own posterior median lattice channel; C′ (bipolar subtraction) subtracts that model\'s median bipolar channel from B, and C′ (all-channel subtraction) additionally subtracts its median extra channel. The two additive estimates coincide for models without extra heat. Values are finite-window log slopes over 400–750 K, not thermal conductivities. Brackets are 95% source-paper bootstrap intervals using 2,000 resamples and seed 2024.")
    definitions=[
        (28,'Component slope medians on identical curves. These population medians must not be subtracted to obtain median paired shifts.','slopes',['A','B','C','Cprime_bip','Cprime_all'],['A','B','C','C′ bipolar','C′ all']),
        (29,'Median per-curve slope differences on identical curves. C minus C′ all measures disagreement between the direct model estimate and additive subtraction; it is not a uniquely identified heat channel.','shifts',['C_B','Cprime_bip_B','Cprime_all_B','C_Cprime_all'],['C − B','C′ bipolar − B','C′ all − B','C − C′ all']),
        (30,'Paired changes from the reference model on identical curves. Each entry is a median of within-curve changes. The change in C − B equals the change in C minus the change in B for each curve, but those population medians need not add. These are descriptive model contrasts, not causal attributions.','changes_from_reference',['B','C','Cprime_bip','Cprime_all','C_B'],['Change in B','Change in C','Change in C′ bipolar','Change in C′ all','Change in C − B'])]
    for number,caption,group,keys,labels in definitions:
        p=si.add_paragraph(f'Supplementary Table {number}. '+caption);p.paragraph_format.keep_with_next=True
        t=si.add_table(rows=1,cols=len(keys)+1);t.style='Table Grid'
        for c,v in zip(t.rows[0].cells,['Model']+labels):c.text=v
        for model in N['models']:
            for c,v in zip(t.add_row().cells,[N['short'][model]]+[fmt(component['models'][model][group][k]) for k in keys]):c.text=v
    for tab in si.tables[-5:]:
        tab.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for ri,row in enumerate(tab.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1
                    if ri==0:p.paragraph_format.keep_with_next=True
                    for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(9);run.bold=ri==0
    # First-citation reference order, without touching mathematical equation nodes.
    paragraphs=list(doc.paragraphs);refindex=next(i for i,p in enumerate(paragraphs) if p.text=='References')
    refs=[p for p in paragraphs[refindex+1:] if re.match(r'^\[\d+\]',p.text)]
    order=[]
    for p in paragraphs[:refindex]:
        for m in re.finditer(r'\[(\d+(?:,\d+)*)\]',p.text):
            for n in map(int,m.group(1).split(',')):
                if 1<=n<=len(refs) and n not in order:order.append(n)
    assert len(order)==len(refs),(order,len(refs))
    mapping={n:i+1 for i,n in enumerate(order)}
    for p in paragraphs[:refindex]:
        for node in p._p.xpath('.//w:t'):
            if node.text:node.text=re.sub(r'\[(\d+(?:,\d+)*)\]',lambda m:'['+','.join(str(mapping[int(x)]) for x in m.group(1).split(','))+']' if all(int(x) in mapping for x in m.group(1).split(',')) else m.group(0),node.text)
    refmap={int(re.match(r'^\[(\d+)\]',p.text).group(1)):p._p for p in refs}
    for p in refs:doc.element.body.remove(p._p)
    for old in order:
        node=refmap[old]
        for text in node.xpath('.//w:t'):
            if text.text:text.text=re.sub(r'^\[\d+\]',f'[{mapping[old]}]',text.text)
        doc.element.body.insert(len(doc.element.body)-1,node)
    for d in [doc,si]:
        d.settings.element.append(OxmlElement('w:doNotAutoCompressPictures'))
        for border in list(d.element.xpath('.//w:pBdr')):border.getparent().remove(border)
        d.core_properties.subject='Unified unpublished project study; reference and extended physical models'
    doc.save(mainpath);si.save(suppath)
    texts=[p.text for p in doc.paragraphs];abstract=texts[texts.index('Abstract')+1]
    (OUT/'structure.json').write_text(json.dumps({'main_tables':len(doc.tables),'main_figures':len(doc.inline_shapes),'supplementary_tables':len(si.tables),'supplementary_figures':len(si.inline_shapes),'abstract_words':len(abstract.split()),'paragraph_words_excluding_tables':sum(len(t.split()) for t in texts),'moved_main_tables':moved,'retained_main_tables':kept,'reference_order':order,'reference_status':'v5 and v6 are unpublished stages of this same project'},indent=2))
    print('Finalised: 7 main figures, 2 main tables, 30 supplementary tables; all original analyses retained.')
if __name__=='__main__':main()
